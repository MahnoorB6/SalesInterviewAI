import os
import re

from pypdf import PdfReader
from docx import Document


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_text_from_pdf(
    file_path: str,
) -> str:
    """Extract text from a PDF resume."""

    reader = PdfReader(
        file_path
    )

    pages = []

    for page in reader.pages:

        text = page.extract_text()

        if text:
            pages.append(
                text.strip()
            )

    return "\n".join(
        pages
    ).strip()


# ============================================================
# DOCX TEXT EXTRACTION
# ============================================================

def extract_text_from_docx(
    file_path: str,
) -> str:
    """Extract text from a DOCX resume."""

    document = Document(
        file_path
    )

    paragraphs = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            paragraphs.append(
                text
            )

    # Also extract text from tables.
    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                cell_text = (
                    cell.text.strip()
                )

                if cell_text:
                    row_text.append(
                        cell_text
                    )

            if row_text:

                paragraphs.append(
                    " | ".join(
                        row_text
                    )
                )

    return "\n".join(
        paragraphs
    ).strip()


# ============================================================
# GENERAL RESUME TEXT EXTRACTION
# ============================================================

def extract_resume_text(
    file_path: str,
) -> str:
    """
    Extract resume text based on file extension.

    Supported:
    - PDF
    - DOCX
    """

    if not file_path:
        raise ValueError(
            "Resume file path is required."
        )

    if not os.path.exists(
        file_path
    ):
        raise FileNotFoundError(
            f"Resume file not found: {file_path}"
        )

    extension = os.path.splitext(
        file_path
    )[1].lower()

    if extension == ".pdf":

        return extract_text_from_pdf(
            file_path
        )

    if extension == ".docx":

        return extract_text_from_docx(
            file_path
        )

    raise ValueError(
        "Unsupported resume format. "
        "Only PDF and DOCX are currently supported."
    )


# ============================================================
# LATEST ROLE DETECTION
# ============================================================

def detect_latest_role(
    resume_text: str,
) -> dict:
    """
    Basic deterministic extraction of the latest role.

    No AI or external API is used.
    """

    result = {
        "latest_role": None,
        "latest_company": None,
        "latest_role_start_date": None,
        "latest_role_end_date": None,
        "latest_role_description": None,
    }

    if not resume_text or not resume_text.strip():
        return result

    lines = [
        line.strip()
        for line in resume_text.splitlines()
        if line.strip()
    ]

    # ========================================================
    # COMMON JOB-TITLE KEYWORDS
    # ========================================================

    role_keywords = [
        "engineer",
        "developer",
        "intern",
        "manager",
        "analyst",
        "executive",
        "assistant",
        "specialist",
        "consultant",
        "designer",
        "administrator",
        "coordinator",
        "representative",
        "associate",
        "officer",
        "lead",
        "scientist",
        "architect",
        "accountant",
        "sales",
        "marketing",
        "director",
        "supervisor",
        "administrator",
        "recruiter",
        "researcher",
    ]

    role_index = None

    # ========================================================
    # FIND FIRST ROLE
    # ========================================================

    for index, line in enumerate(
        lines
    ):

        lower_line = line.lower()

        if any(
            keyword in lower_line
            for keyword in role_keywords
        ):

            role_index = index

            break

    # ========================================================
    # ROLE + COMPANY
    # ========================================================

    if role_index is not None:

        result[
            "latest_role"
        ] = lines[role_index]

        if (
            role_index + 1
            < len(lines)
        ):

            possible_company = (
                lines[
                    role_index + 1
                ]
            )

            possible_company_lower = (
                possible_company.lower()
            )

            if not any(
                keyword in possible_company_lower
                for keyword in role_keywords
            ):

                result[
                    "latest_company"
                ] = possible_company

        # ====================================================
        # ROLE DESCRIPTION
        # ====================================================

        description_lines = lines[
            role_index + 2:
            role_index + 7
        ]

        if description_lines:

            result[
                "latest_role_description"
            ] = " ".join(
                description_lines
            )

    # ========================================================
    # DATE RANGE
    # ========================================================

    date_pattern = re.compile(
        r"""
        (
            (?:
                Jan|Feb|Mar|Apr|May|Jun|
                Jul|Aug|Sep|Oct|Nov|Dec
            )
            [a-z]*
            \s+
            (?:19|20)\d{2}
        )
        \s*
        (?:-|–|—|to)
        \s*
        (
            (?:
                (?:
                    Jan|Feb|Mar|Apr|May|Jun|
                    Jul|Aug|Sep|Oct|Nov|Dec
                )
                [a-z]*
                \s+
                (?:19|20)\d{2}
            )
            |
            Present
            |
            Current
            |
            Now
        )
        """,
        re.IGNORECASE |
        re.VERBOSE,
    )

    date_match = (
        date_pattern.search(
            resume_text
        )
    )

    if date_match:

        result[
            "latest_role_start_date"
        ] = date_match.group(1)

        result[
            "latest_role_end_date"
        ] = date_match.group(2)

    return result